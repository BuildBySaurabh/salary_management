import calendar
import csv
from datetime import date, datetime
from io import BytesIO
from django.contrib.auth import get_user_model
from django.db import IntegrityError, transaction
from django.db.models import Avg, Count, Q, Sum
from django.http import HttpResponse, StreamingHttpResponse
from django.utils.dateparse import parse_date
from openpyxl import Workbook, load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from rest_framework import status, viewsets
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response
from rest_framework_simplejwt.tokens import RefreshToken

from .filters import EmployeeFilter
from .models import Employee, SalaryRecord, SalaryRevision
from .pagination import EmployeePagination
from .permissions import IsHRManager
from .payroll import calculate_salary
from .serializers import EmployeeImportSerializer, EmployeeSerializer, MeSerializer, RegisterSerializer, SalaryRecordSerializer, SalaryRevisionSerializer

User = get_user_model()


@api_view(["POST"])
@permission_classes([AllowAny])
def register(request):
    serializer = RegisterSerializer(data=request.data)
    serializer.is_valid(raise_exception=True)
    try:
        with transaction.atomic(): user = serializer.save()
    except IntegrityError:
        return Response({"detail": "An account with this email already exists. Please sign in."}, status=400)
    refresh = RefreshToken.for_user(user)
    return Response({"user": MeSerializer(user).data, "access": str(refresh.access_token), "refresh": str(refresh)}, status=201)


@api_view(["GET"])
def me(request):
    return Response(MeSerializer(request.user).data)


class EmployeeViewSet(viewsets.ModelViewSet):
    queryset = Employee.objects.all()
    serializer_class = EmployeeSerializer
    permission_classes = [IsHRManager]
    pagination_class = EmployeePagination
    filterset_class = EmployeeFilter
    ordering_fields = ["employee_id", "first_name", "department", "country", "annual_salary", "hire_date", "status"]
    ordering = ["employee_id"]

    def get_queryset(self):
        qs = super().get_queryset()
        search = self.request.query_params.get("search", "").strip()
        if search:
            terms = search.split()
            query = Q(employee_id__icontains=search) | Q(first_name__icontains=search) | Q(last_name__icontains=search) | Q(email__icontains=search) | Q(phone__icontains=search) | Q(job_title__icontains=search) | Q(department__icontains=search) | Q(country__icontains=search)
            if len(terms) >= 2:
                query |= Q(first_name__icontains=terms[0], last_name__icontains=terms[-1])
            qs = qs.filter(query)
        return qs

    @action(detail=False, methods=["post"], url_path="import")
    def import_excel(self, request):
        serializer = EmployeeImportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        workbook = load_workbook(serializer.validated_data["file"], read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows: return Response({"detail": "The spreadsheet is empty."}, status=400)
        headers = [str(v).strip().lower() if v is not None else "" for v in rows[0]]
        required = ["employee_id", "first_name", "last_name", "email", "job_title", "department", "country", "currency", "salary", "hire_date"]
        missing = [field for field in required if field not in headers]
        if missing: return Response({"detail": f"Missing columns: {', '.join(missing)}"}, status=400)
        index = {name: i for i, name in enumerate(headers)}
        created = updated = 0; errors = []
        for line_number, row in enumerate(rows[1:], start=2):
            if not any(value is not None for value in row): continue
            try:
                employee_id = str(row[index["employee_id"]]).strip().upper()
                hire_date = row[index["hire_date"]]
                if hasattr(hire_date, "date"): hire_date = hire_date.date()
                elif isinstance(hire_date, str): hire_date = datetime.strptime(hire_date, "%Y-%m-%d").date()
                payload = {
                    "employee_id": employee_id, "first_name": str(row[index["first_name"]]).strip(), "last_name": str(row[index["last_name"]]).strip(),
                    "email": str(row[index["email"]]).strip().lower(), "phone": str(row[index["phone"]]).strip() if "phone" in index and row[index["phone"]] else "",
                    "job_title": str(row[index["job_title"]]).strip(), "department": str(row[index["department"]]).strip(), "country": str(row[index["country"]]).strip(),
                    "currency": str(row[index["currency"]]).strip().upper(), "salary": row[index["salary"]], "hire_date": hire_date,
                    "status": str(row[index["status"]]).strip().lower() if "status" in index and row[index["status"]] else "active", "source": "import",
                    "bank_name": str(row[index["bank_name"]]).strip() if "bank_name" in index and row[index["bank_name"]] else "",
                    "bank_branch": str(row[index["bank_branch"]]).strip() if "bank_branch" in index and row[index["bank_branch"]] else "",
                    "bank_account_last4": str(row[index["bank_account_last4"]]).strip()[-4:] if "bank_account_last4" in index and row[index["bank_account_last4"]] else "",
                }
                obj = Employee.objects.filter(employee_id=employee_id).first()
                ser = EmployeeSerializer(obj, data=payload, partial=True) if obj else EmployeeSerializer(data=payload)
                ser.is_valid(raise_exception=True); ser.save()
                updated += 1 if obj else 0; created += 0 if obj else 1
            except Exception as exc:
                errors.append({"row": line_number, "error": str(exc)})
        workbook.close()
        return Response({"created": created, "updated": updated, "errors": errors[:25], "error_count": len(errors)})

    @action(detail=False, methods=["get"], url_path="export")
    def export_csv(self, request):
        qs = self.filter_queryset(self.get_queryset())
        response = StreamingHttpResponse(content_type="text/csv")
        response["Content-Disposition"] = 'attachment; filename="acme-employees.csv"'
        writer = csv.writer(response)
        writer.writerow(["Employee ID", "Name", "Email", "Phone", "Job Title", "Department", "Country", "Currency", "Annual CTC", "Hire Date", "Status", "Bank Name", "Branch", "Account Last 4"])
        for employee in qs.iterator():
            writer.writerow([employee.employee_id, employee.full_name, employee.email, employee.phone, employee.job_title, employee.department, employee.country, employee.currency, employee.annual_salary, employee.hire_date, employee.status, employee.bank_name, employee.bank_branch, employee.bank_account_last4])
        return response

    @action(detail=True, methods=["get"], url_path="salary-history")
    def salary_history(self, request, pk=None):
        employee = self.get_object()
        return Response(SalaryRevisionSerializer(employee.salary_revisions.all(), many=True).data)

    @action(detail=True, methods=["post"], url_path="generate-salary-slip")
    def generate_salary_slip(self, request, pk=None):
        employee = self.get_object()
        month = request.data.get("month") or request.query_params.get("month")
        if not month:
            today = date.today(); month = f"{today.year:04d}-{today.month:02d}"
        try: year, month_number = [int(x) for x in month.split("-")]
        except Exception: return Response({"detail": "Month must use YYYY-MM format."}, status=400)
        if month_number < 1 or month_number > 12: return Response({"detail": "Invalid month."}, status=400)
        pay_month = date(year, month_number, 1)
        record, _ = create_or_update_salary_record(employee, pay_month)
        return salary_slip_response(record)


class SalaryRecordViewSet(viewsets.ModelViewSet):
    queryset = SalaryRecord.objects.select_related("employee").all()
    serializer_class = SalaryRecordSerializer
    permission_classes = [IsHRManager]
    pagination_class = EmployeePagination
    http_method_names = ["get", "post", "patch", "head", "options"]

    def perform_update(self, serializer):
        record = serializer.save()
        if record.payment_status == SalaryRecord.PaymentStatus.PAID and not record.paid_date:
            record.paid_date = date.today()
            record.save(update_fields=["paid_date", "updated_at"])

    def get_queryset(self):
        qs = super().get_queryset()
        month = self.request.query_params.get("month")
        search = self.request.query_params.get("search", "").strip()
        payment_status = self.request.query_params.get("payment_status")
        if month:
            try: qs = qs.filter(pay_month=date.fromisoformat(f"{month}-01"))
            except ValueError: pass
        if payment_status: qs = qs.filter(payment_status=payment_status)
        if search:
            qs = qs.filter(Q(employee_name_snapshot__icontains=search) | Q(employee_id_snapshot__icontains=search) | Q(department__icontains=search))
        return qs

    @action(detail=True, methods=["get"], url_path="slip")
    def slip(self, request, pk=None):
        return salary_slip_response(self.get_object())

    @action(detail=False, methods=["post"], url_path="generate")
    def generate(self, request):
        month = request.data.get("month")
        if not month: return Response({"detail": "Select a payroll month."}, status=400)
        try: year, month_number = [int(x) for x in month.split("-")]
        except Exception: return Response({"detail": "Month must use YYYY-MM format."}, status=400)
        if month_number not in range(1, 13): return Response({"detail": "Invalid month."}, status=400)
        pay_month = date(year, month_number, 1)
        employees = Employee.objects.filter(status__in=[Employee.Status.ACTIVE, Employee.Status.ON_LEAVE]).order_by("employee_id")
        employee_ids = request.data.get("employee_ids") or []
        if employee_ids: employees = employees.filter(id__in=employee_ids)
        created = updated = 0
        for employee in employees.iterator():
            _, was_created = create_or_update_salary_record(employee, pay_month)
            created += int(was_created); updated += int(not was_created)
        return Response({"month": month, "created": created, "updated": updated, "count": created + updated})

    @action(detail=False, methods=["post"], url_path="import")
    def import_excel(self, request):
        file = request.FILES.get("file")
        if not file: return Response({"detail": "Please attach an Excel salary sheet."}, status=400)
        workbook = load_workbook(file, read_only=True, data_only=True); sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True)); workbook.close()
        if not rows: return Response({"detail": "The spreadsheet is empty."}, status=400)
        headers = [str(v).strip().lower() if v is not None else "" for v in rows[0]]; idx = {h: i for i, h in enumerate(headers)}
        required = ["employee_id", "pay_month", "payment_status"]
        missing = [x for x in required if x not in idx]
        if missing: return Response({"detail": f"Missing columns: {', '.join(missing)}"}, status=400)
        created = updated = 0; errors = []
        for line, row in enumerate(rows[1:], 2):
            if not any(v is not None and str(v) != "" for v in row): continue
            try:
                employee = Employee.objects.get(employee_id=str(row[idx["employee_id"]]).strip().upper())
                raw_month = row[idx["pay_month"]]
                if hasattr(raw_month, "date"): pay_month = raw_month.date().replace(day=1)
                else: pay_month = date.fromisoformat(str(raw_month)[:7] + "-01")
                record, was_created = create_or_update_salary_record(employee, pay_month, preserve_import=True, row=row, idx=idx)
                created += int(was_created); updated += int(not was_created)
            except Exception as exc:
                errors.append({"row": line, "error": str(exc)})
        return Response({"created": created, "updated": updated, "error_count": len(errors), "errors": errors[:25]})

    @action(detail=False, methods=["get"], url_path="export")
    def export_excel(self, request):
        records = self.filter_queryset(self.get_queryset()).order_by("employee_id_snapshot")
        wb = Workbook(); ws = wb.active; ws.title = "Salary Sheet"
        headers = ["Employee ID", "Employee Name", "Pay Month", "CTC", "Department", "Currency", "Bank Name", "Branch", "Account Last 4", "Working Days", "Paid Days", "Basic", "HRA", "Allowances", "Bonus", "Overtime", "Gross Pay", "PF", "Professional Tax", "Income Tax", "Other Deductions", "Total Deductions", "Net Pay", "Payment Status", "Paid Date"]
        ws.append(headers)
        for r in records.iterator():
            ws.append([r.employee_id_snapshot, r.employee_name_snapshot, r.pay_month, r.ctc, r.department, r.currency, r.bank_name, r.bank_branch, r.bank_account_last4, r.working_days, r.paid_days, r.basic_pay, r.hra, r.allowances, r.bonus, r.overtime, r.gross_pay, r.pf_amount, r.professional_tax, r.income_tax, r.other_deductions, r.total_deductions, r.net_pay, r.payment_status, r.paid_date])
        for cell in ws[1]: cell.font = cell.font.copy(bold=True)
        output = BytesIO(); wb.save(output); output.seek(0)
        response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        response["Content-Disposition"] = 'attachment; filename="acme-salary-sheet.xlsx"'
        return response


def create_or_update_salary_record(employee, pay_month, preserve_import=False, row=None, idx=None):
    calculated = calculate_salary(employee, pay_month.year, pay_month.month)
    defaults = {
        **calculated,
        "employee_id_snapshot": employee.employee_id,
        "employee_name_snapshot": employee.full_name,
        "department": employee.department,
        "currency": employee.currency,
        "bank_name": employee.bank_name,
        "bank_branch": employee.bank_branch,
        "bank_account_last4": employee.bank_account_last4,
        "payment_status": SalaryRecord.PaymentStatus.UNPAID,
    }
    if preserve_import and row is not None:
        def value(name, default): return row[idx[name]] if name in idx and row[idx[name]] not in (None, "") else default
        defaults["payment_status"] = str(value("payment_status", defaults["payment_status"])).lower()
        if "paid_date" in idx and value("paid_date", None):
            raw = value("paid_date", None); defaults["paid_date"] = raw.date() if hasattr(raw, "date") else parse_date(str(raw))
        for field in ["paid_days", "basic_pay", "hra", "allowances", "bonus", "overtime", "gross_pay", "pf_amount", "professional_tax", "income_tax", "other_deductions", "total_deductions", "net_pay"]:
            if field in idx and value(field, None) not in (None, ""): defaults[field] = value(field, defaults[field])
    record, created = SalaryRecord.objects.update_or_create(employee=employee, pay_month=pay_month, defaults=defaults)
    return record, created


def salary_slip_response(record):
    styles = getSampleStyleSheet()
    title = ParagraphStyle("SlipTitle", parent=styles["Title"], fontSize=18, leading=22, spaceAfter=8)
    small = ParagraphStyle("Small", parent=styles["BodyText"], fontSize=8.5, leading=11)
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=A4, rightMargin=16*mm, leftMargin=16*mm, topMargin=14*mm, bottomMargin=14*mm)
    currency = record.currency
    money_rows = [
        ["Basic Pay", f"{currency} {record.basic_pay:,.2f}"], ["HRA", f"{currency} {record.hra:,.2f}"],
        ["Allowances", f"{currency} {record.allowances:,.2f}"], ["Bonus", f"{currency} {record.bonus:,.2f}"],
        ["Overtime", f"{currency} {record.overtime:,.2f}"], ["Gross Pay", f"{currency} {record.gross_pay:,.2f}"],
        ["PF / Retirement", f"{currency} {record.pf_amount:,.2f}"], ["Professional Tax", f"{currency} {record.professional_tax:,.2f}"],
        ["Income Tax / Withholding", f"{currency} {record.income_tax:,.2f}"], ["Other Deductions", f"{currency} {record.other_deductions:,.2f}"],
        ["Total Deductions", f"{currency} {record.total_deductions:,.2f}"], ["Net Pay", f"{currency} {record.net_pay:,.2f}"],
    ]
    story = [Paragraph("ACME Salary Manager", title), Paragraph("Employee Salary Slip", styles["Heading2"]), Spacer(1, 5)]
    story.append(Table([
        ["Employee", record.employee_name_snapshot], ["Employee ID", record.employee_id_snapshot], ["Department", record.department],
        ["Pay Month", record.pay_month.strftime("%B %Y")], ["CTC", f"{currency} {record.ctc:,.2f}"],
        ["Bank", f"{record.bank_name or '—'} · {record.bank_branch or '—'} · ****{record.bank_account_last4 or '—'}"],
        ["Working / Paid Days", f"{record.working_days} / {record.paid_days}"], ["Payment Status", record.get_payment_status_display()],
    ], colWidths=[45*mm, 125*mm], style=[("BACKGROUND", (0,0),(0,-1), colors.HexColor("#f3f6fa")), ("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#dbe3ee")), ("FONTNAME",(0,0),(0,-1),"Helvetica-Bold"), ("FONTSIZE",(0,0),(-1,-1),9), ("VALIGN",(0,0),(-1,-1),"TOP")]))
    story += [Spacer(1, 10), Paragraph("Payroll calculation", styles["Heading3"])]
    story.append(Table(money_rows, colWidths=[105*mm, 65*mm], style=[("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#dbe3ee")), ("FONTNAME",(0,5),(1,5),"Helvetica-Bold"), ("FONTNAME",(0,-1),(1,-1),"Helvetica-Bold"), ("BACKGROUND",(0,5),(1,5),colors.HexColor("#eef5ff")), ("BACKGROUND",(0,-1),(1,-1),colors.HexColor("#eaf8f2")), ("ALIGN",(1,0),(1,-1),"RIGHT"), ("FONTSIZE",(0,0),(-1,-1),9)]))
    story += [Spacer(1, 12), Paragraph("Demo payroll note: salary components use a configurable standard calculation for this assessment. Statutory tax, social-security and local payroll rules should be configured by the HR/payroll team before production use.", small)]
    doc.build(story); output.seek(0)
    response = HttpResponse(output.getvalue(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="salary-slip-{record.employee_id_snapshot}-{record.pay_month:%Y-%m}.pdf"'
    return response


@api_view(["GET"])
def dashboard_summary(request):
    if not request.user.is_authenticated: return Response({"detail": "Authentication credentials were not provided."}, status=401)
    total = Employee.objects.count(); countries = Employee.objects.values("country").distinct().count(); departments = Employee.objects.values("department").distinct().count(); currencies = Employee.objects.values("currency").distinct().count()
    by_country = list(Employee.objects.values("country").annotate(value=Count("id")).order_by("-value", "country")[:12])
    by_department = list(Employee.objects.values("department").annotate(value=Count("id")).order_by("-value", "department"))
    by_currency = list(Employee.objects.values("currency").annotate(employees=Count("id"), average_salary=Avg("annual_salary"), total_payroll=Sum("annual_salary")).order_by("-employees"))
    return Response({"kpis": {"total_employees": total, "countries": countries, "departments": departments, "currencies": currencies, "active": Employee.objects.filter(status="active").count(), "on_leave": Employee.objects.filter(status="on_leave").count(), "inactive": Employee.objects.filter(status="inactive").count(), "salary_records": SalaryRecord.objects.count()}, "by_country": by_country, "by_department": by_department, "by_currency": [{**x, "average_salary": float(x["average_salary"] or 0), "total_payroll": float(x["total_payroll"] or 0)} for x in by_currency]})


@api_view(["GET"])
def report_summary(request):
    if not request.user.is_authenticated: return Response({"detail": "Authentication credentials were not provided."}, status=401)
    department = list(Employee.objects.values("department").annotate(employees=Count("id"), avg_salary=Avg("annual_salary"), payroll=Sum("annual_salary")).order_by("-payroll"))
    country = list(Employee.objects.values("country").annotate(employees=Count("id"), avg_salary=Avg("annual_salary")).order_by("-employees"))
    monthly = list(SalaryRecord.objects.values("pay_month", "currency").annotate(gross=Sum("gross_pay"), net=Sum("net_pay"), employees=Count("id")).order_by("-pay_month")[:24])
    for item in monthly: item["pay_month"] = item["pay_month"].isoformat(); item["gross"] = float(item["gross"] or 0); item["net"] = float(item["net"] or 0)
    return Response({"department": [{**x, "avg_salary": float(x["avg_salary"] or 0), "payroll": float(x["payroll"] or 0)} for x in department], "country": [{**x, "avg_salary": float(x["avg_salary"] or 0)} for x in country], "monthly": monthly})
